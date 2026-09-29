import json, os
HERE = os.path.dirname(os.path.abspath(__file__))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

d = json.load(open(os.path.join(HERE, "..", "js", "checklist.json")))
teams = {t["id"]: t["name"] for t in d["teams"]}
wb = Workbook()
F = "Arial"
HEAD = PatternFill("solid", fgColor="1F4E4A")
BAND = PatternFill("solid", fgColor="EEF4F3")
INPUT = PatternFill("solid", fgColor="FFF6D5")
thin = Side(style="thin", color="C9D6D4")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
SEVCOL = {"critical": "B3261E", "major": "B26A00", "minor": "3D6B63"}


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = HEAD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def body(ws, r, vals, band=False):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name=F, size=10)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BOX
        if band:
            c.fill = BAND


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(name=F, bold=True, size=14, color="1F4E4A")
    ws["A2"] = sub
    ws["A2"].font = Font(name=F, size=10, italic=True, color="555555")


# ---- README
ws = wb.active
ws.title = "Read me"
title(ws, "KVD Windpark: flat snagging checklist", "RCC framed towers, plaster and paint, aluminium / uPVC windows, wooden internal doors. Phase 1: towers T1, T2, T3, T6, T7 (T4, T5 are Phase 2). 24 floors, 2BHK and 3BHK.")
lines = [
    ("Checklist", "Every check item. One row per item, grouped by building element (ceiling, walls, floor, doors, windows, services). ID is what the app stores."),
    ("Rooms by flat type", "Which rooms each flat type has and which element groups apply to each room. Check counts are live formulas."),
    ("Teams", "Default responsible team for each item. Edit team names here and in the app's Setup screen."),
    ("Severity", "What Critical, Major and Minor mean and target fix times."),
    ("Handover", "Items recorded at handover: gate conditions, meter readings, keys, documents, demonstrations."),
    ("Apartment list", "Fill this in (yellow cells) and paste into the app's Setup screen to load the actual flats. Row 5 is an example."),
    ("", ""),
    ("Criteria", "Tolerances marked 'as per spec' or 'typically' are working defaults. Replace them with the figures in the KVD contract specifications before issue."),
    ("Workflow", "Open (raised by QC, assigned to team) > Fixed (team marks done with photo) > Closed (QC verifies). QC can reopen. Handover is allowed only when no Critical or Major snag is open."),
]
for i, (a, b) in enumerate(lines, 4):
    ws.cell(row=i, column=1, value=a).font = Font(name=F, bold=True, size=10)
    c = ws.cell(row=i, column=2, value=b)
    c.font = Font(name=F, size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 110

# ---- Checklist
ws = wb.create_sheet("Checklist")
title(ws, "Checklist master", f"{len(d['items'])} check items. Applies per room as mapped on 'Rooms by flat type'.")
header(ws, 4, ["ID", "Element group", "Check", "Acceptance criterion / method", "Default team", "Severity", "Applies to rooms"],
       [10, 26, 60, 60, 24, 11, 36])
kind_by_group = {}
for k, v in d["kinds"].items():
    for g in v["groups"]:
        kind_by_group.setdefault(g, []).append(v["label"])
r = 5
for n, it in enumerate(d["items"]):
    body(ws, r, [it["id"], d["groups"][it["g"]], it["t"], it["c"], teams[it["team"]], it["sev"].capitalize(),
                 ", ".join(kind_by_group.get(it["g"], []))], band=(n % 2 == 1))
    ws.cell(row=r, column=6).font = Font(name=F, size=10, bold=True, color=SEVCOL[it["sev"]])
    r += 1
ws.auto_filter.ref = f"A4:G{r-1}"
last_item_row = r - 1

# ---- Rooms by flat type
ws = wb.create_sheet("Rooms by flat type")
title(ws, "Rooms by flat type", "Check count = number of checklist rows whose element group is listed for the room (live formula).")
header(ws, 4, ["Flat type", "Room key", "Room", "Room kind", "Element groups", "Checks"], [11, 10, 26, 18, 60, 9])
r = 5
code_of = {v: k for k, v in d["groups"].items()}
for ft, rooms in d["flatTypes"].items():
    start = r
    for n, rm in enumerate(rooms):
        groups = d["kinds"][rm["kind"]]["groups"]
        body(ws, r, [ft, rm["key"], rm["label"], d["kinds"][rm["kind"]]["label"], ", ".join(groups), None], band=(n % 2 == 1))
        f = "+".join(f'COUNTIF(Checklist!$B$5:$B${last_item_row},"{d["groups"][g]}")' for g in groups)
        ws.cell(row=r, column=6, value="=" + f)
        r += 1
    ws.cell(row=r, column=5, value=f"Total checks per {ft} flat").font = Font(name=F, bold=True, size=10)
    c = ws.cell(row=r, column=6, value=f"=SUM(F{start}:F{r-1})")
    c.font = Font(name=F, bold=True, size=10)
    r += 2

# ---- Teams
ws = wb.create_sheet("Teams")
title(ws, "Teams", "Default owner for snags raised against each item. Counts are live formulas.")
header(ws, 4, ["Team ID", "Team", "Checklist items owned", "of which Critical"], [16, 30, 22, 18])
for n, t in enumerate(d["teams"]):
    rr = 5 + n
    body(ws, rr, [t["id"], t["name"], None, None], band=(n % 2 == 1))
    ws.cell(row=rr, column=3, value=f'=COUNTIF(Checklist!$E$5:$E${last_item_row},B{rr})')
    ws.cell(row=rr, column=4, value=f'=COUNTIFS(Checklist!$E$5:$E${last_item_row},B{rr},Checklist!$F$5:$F${last_item_row},"Critical")')

# ---- Severity
ws = wb.create_sheet("Severity")
title(ws, "Severity", "Used on every snag.")
header(ws, 4, ["Severity", "Meaning", "Items in checklist"], [12, 90, 18])
for n, s in enumerate(d["severity"]):
    rr = 5 + n
    body(ws, rr, [s["name"], s["desc"], None])
    ws.cell(row=rr, column=1).font = Font(name=F, size=10, bold=True, color=SEVCOL[s["id"]])
    ws.cell(row=rr, column=3, value=f'=COUNTIF(Checklist!$F$5:$F${last_item_row},A{rr})')

# ---- Handover
ws = wb.create_sheet("Handover")
title(ws, "Handover record", "Captured in the app's Handover tab for each flat, with customer and company signatures.")
header(ws, 4, ["Section", "Item", "Record"], [26, 70, 30])
sec = {"gate": "Before handover", "meters": "Meter readings", "keys": "Keys handed (count)", "docs": "Documents handed", "demo": "Shown to customer"}
rr = 5
body(ws, rr, ["Before handover", "No Critical or Major snag open (checked automatically)", "System"]); rr += 1
body(ws, rr, ["Before handover", "Inspection 100% complete (checked automatically)", "System"]); rr += 1
for key in ["gate", "meters", "keys", "docs", "demo"]:
    for it in d["handover"][key]:
        rec = {"meters": "Reading", "keys": "Number", "gate": "Tick", "docs": "Tick", "demo": "Tick"}[key]
        body(ws, rr, [sec[key], it["label"], rec]); rr += 1
for lab in ["Customer name", "Customer phone", "Customer remarks", "Minor snags accepted with fix date", "Customer signature", "Company representative signature", "Handover date"]:
    body(ws, rr, ["Acceptance", lab, "Entry"]); rr += 1

# ---- Apartment list
ws = wb.create_sheet("Apartment list")
title(ws, "Apartment list", "Fill yellow cells from row 6 down. Copy columns A-E (with header row) and paste into the app: Setup > Apartment list.")
header(ws, 4, ["Tower", "Floor", "Flat no.", "Type", "Carpet area (sq ft)"], [12, 8, 10, 10, 20])
body(ws, 5, ["T1", 12, 1203, "3BHK", 1450])
for c in range(1, 6):
    ws.cell(row=5, column=c).font = Font(name=F, size=10, italic=True, color="777777")
ws.cell(row=5, column=6, value="Example row: delete before pasting").font = Font(name=F, size=9, italic=True, color="777777")
for rr in range(6, 206):
    for c in range(1, 6):
        cell = ws.cell(row=rr, column=c)
        cell.fill = INPUT
        cell.border = BOX
        cell.font = Font(name=F, size=10)
dv = DataValidation(type="list", formula1='"2BHK,3BHK"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("D6:D2000")

out = os.path.join(HERE, "..", "docs", "KVD_Windpark_Snag_Checklist.xlsx")
wb.save(out)
print(out)
